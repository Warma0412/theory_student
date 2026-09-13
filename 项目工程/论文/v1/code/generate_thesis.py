#!/usr/bin/env python3
"""Generate the v1 thesis from audited result files."""

from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_panel import (
    CANDIDATE_FEATURES,
    EXCLUDED_FEATURES,
    FEATURES,
    FEATURE_LABELS_ZH,
    V1_DIR,
)


RESULTS = V1_DIR / "results"
FIGURES = V1_DIR / "figures"
TEMPLATE = Path(__file__).with_name("thesis_template.md")


def load_json(name: str):
    return json.loads((RESULTS / name).read_text(encoding="utf-8"))


def fmt(value, digits: int = 3) -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return "-"
    if isinstance(value, (int, np.integer)):
        return f"{value:,}"
    return f"{float(value):,.{digits}f}"


def markdown_table(
    frame: pd.DataFrame,
    columns: list[str],
    headers: list[str],
    formatter: dict | None = None,
) -> str:
    formatter = formatter or {}
    rows = [
        "| " + " | ".join(headers) + " |",
        "|" + "|".join(["---"] * len(headers)) + "|",
    ]
    for _, row in frame[columns].iterrows():
        cells = []
        for column in columns:
            value = row[column]
            if column in formatter:
                value = formatter[column](value)
            elif isinstance(value, (float, np.floating)):
                value = fmt(value)
            cells.append(str(value).replace("|", "\\|").replace("\n", " "))
        rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def zh_list(features: list[str]) -> str:
    if not features:
        return "无"
    return "、".join(FEATURE_LABELS_ZH.get(feature, feature) for feature in features)


def scenario_name(target: str, generator: str, model: str) -> str:
    target_names = {
        "log_gmv_next_month": "次月GMV",
        "log_gmv": "同月GMV",
        "log_aov_next_month": "次月客单价",
    }
    generator_names = {
        "copula": "Copula-MVR",
        "gaussian": "原始高斯-MVR",
        "copula_trimmed_period": "Copula-MVR（截尾期）",
        "copula_ai_xgboost": "Copula-MVR",
        "copula_pretest": "Copula-MVR（测试前）",
    }
    model_names = {"lasso": "Lasso", "xgboost": "XGBoost"}
    return (
        f"{target_names.get(target, target)}；"
        f"{generator_names.get(generator, generator)}；"
        f"{model_names.get(model, model)}"
    )


def build_context() -> dict[str, str]:
    audit = load_json("data_audit.json")
    repro = load_json("reproducibility.json")
    diagnostics = load_json("knockoff_diagnostics.json")
    simulation = load_json("simulation_calibration.json")
    primary = pd.read_csv(RESULTS / "knockoff_primary.csv")
    scenarios = pd.read_csv(RESULTS / "knockoff_all_scenarios.csv")
    descriptive = pd.read_csv(RESULTS / "descriptive_statistics.csv")
    metrics = pd.read_csv(RESULTS / "predictive_model_metrics.csv")
    shap = pd.read_csv(RESULTS / "xgboost_shap_importance.csv")
    fixed = pd.read_csv(RESULTS / "two_way_fixed_effects.csv")
    monthly = pd.read_csv(RESULTS / "monthly_summary.csv")
    dictionary = pd.read_csv(RESULTS / "feature_dictionary.csv")

    p20 = primary.loc[primary["q"].eq(0.20)].sort_values(
        ["selection_frequency", "mean_w"], ascending=False
    )
    strict = p20.loc[p20["selected_ebh"].astype(bool), "feature"].tolist()
    stable = p20.loc[p20["selection_frequency"].ge(0.90), "feature"].tolist()
    q10 = primary.loc[
        primary["q"].eq(0.10) & primary["selected_ebh"].astype(bool), "feature"
    ].tolist()
    q30 = primary.loc[
        primary["q"].eq(0.30) & primary["selected_ebh"].astype(bool), "feature"
    ].tolist()
    pretest = repro["pretest_selected_features"]
    top_shap = shap.head(10)["feature"].tolist()
    core = [feature for feature in strict if feature in top_shap]
    secondary = [feature for feature in strict if feature not in core]
    if not core:
        core = [feature for feature in stable if feature in top_shap]
        secondary = [feature for feature in stable if feature not in core]

    best = metrics.sort_values("rmse_log").iloc[0]
    naive = metrics.loc[metrics["model"].eq("Naive-current-GMV")].iloc[0]
    improvement = (naive["rmse_log"] - best["rmse_log"]) / naive["rmse_log"]
    diag_copula = next(item for item in diagnostics if item["generator"] == "copula")
    diag_gaussian = next(item for item in diagnostics if item["generator"] == "gaussian")
    significant = fixed.loc[fixed["p_value"].lt(0.05)].sort_values("p_value")
    peak = monthly.loc[monthly["gmv"].idxmax()]

    source_table = pd.DataFrame(
        [
            ["订单", "olist_orders_dataset.csv", audit["source_rows"]["orders"], "状态、时间戳"],
            ["商品明细", "olist_order_items_dataset.csv", audit["source_rows"]["items"], "卖家、商品、价格、运费"],
            ["支付", "olist_order_payments_dataset.csv", audit["source_rows"]["payments"], "方式、分期、金额"],
            ["评价", "olist_order_reviews_dataset.csv", audit["source_rows"]["reviews"], "评分"],
            ["商品", "olist_products_dataset.csv", audit["source_rows"]["products"], "品类、物理与内容属性"],
            ["卖家", "olist_sellers_dataset.csv", audit["source_rows"]["sellers"], "卖家位置"],
            ["顾客", "olist_customers_dataset.csv", audit["source_rows"]["customers"], "顾客唯一标识与位置"],
            ["地理", "olist_geolocation_dataset.csv", audit["source_rows"]["geolocation"], "邮编前缀经纬度"],
            ["营销线索", "olist_marketing_qualified_leads_dataset.csv", audit["source_rows"]["mql"], "线索来源"],
            ["成交线索", "olist_closed_deals_dataset.csv", audit["source_rows"]["closed_deals"], "线索与卖家映射"],
        ],
        columns=["数据域", "文件", "行数", "用途"],
    )

    dictionary["编号"] = [f"X{i + 1}" for i in range(len(dictionary))]
    dictionary["状态"] = dictionary["analysis_included"].map(
        lambda value: "纳入分析" if bool(value) else "零方差排除"
    )

    key_features = [
        "avg_price",
        "avg_freight",
        "avg_review_score",
        "avg_delivery_days",
        "avg_ship_days",
        "item_count",
        "order_count",
        "category_count",
        "unique_customer_count",
        "buyer_state_diversity",
        "interstate_ratio",
        "avg_distance_km",
    ]
    key_desc = descriptive.loc[descriptive["feature"].isin(key_features)].copy()
    key_desc["missing_pct"] = key_desc["missing_rate"].map(lambda value: f"{value:.2%}")

    primary_table = p20.head(20).copy()
    primary_table["frequency_fmt"] = primary_table["selection_frequency"].map(
        lambda value: f"{value:.1%}"
    )
    primary_table["strict_fmt"] = primary_table["selected_ebh"].map(
        lambda value: "通过" if bool(value) else "未通过"
    )

    metric_table = metrics.copy()
    metric_table["wape_pct"] = metric_table["wape_raw"].map(lambda value: f"{value:.1%}")

    shap_table = shap.head(15).copy()
    fixed_view = fixed.sort_values("p_value").head(15).copy()

    scenario_rows = []
    for keys, group in scenarios.groupby(
        ["target", "generator", "stat_model", "q", "alpha_kn"], dropna=False
    ):
        selected = group.loc[group["selected_ebh"].astype(bool), "label_zh"].tolist()
        stable_group = group.loc[group["selection_frequency"].ge(0.90), "label_zh"].tolist()
        scenario_rows.append(
            {
                "设定": scenario_name(keys[0], keys[1], keys[2]),
                "alpha_eBH": keys[3],
                "alpha_kn": keys[4],
                "严格入选数": len(selected),
                "严格入选变量": "、".join(selected) if selected else "无",
                "稳定候选数": len(stable_group),
            }
        )
    scenario_table = pd.DataFrame(scenario_rows)

    sensitivity_table = pd.DataFrame(
        [
            ["0.10", f"{0.05:.2f}", len(q10), zh_list(q10)],
            ["0.20（主分析）", f"{0.10:.2f}", len(strict), zh_list(strict)],
            ["0.30", f"{0.15:.2f}", len(q30), zh_list(q30)],
        ],
        columns=["alpha_eBH", "alpha_kn", "严格入选数", "严格入选变量"],
    )

    top_corr = descriptive.reindex(
        descriptive["spearman_next_gmv"].abs().sort_values(ascending=False).index
    ).head(10)
    appendix_desc = descriptive.copy()
    appendix_desc["missing_pct"] = appendix_desc["missing_rate"].map(
        lambda value: f"{value:.2%}"
    )

    def select_scenario(generator: str) -> list[str]:
        subset = scenarios.loc[
            scenarios["generator"].eq(generator)
            & scenarios["q"].eq(0.20)
            & scenarios["selected_ebh"].astype(bool),
            "feature",
        ]
        return subset.tolist()

    context = {
        "PANEL_ROWS": f"{audit['panel_rows']:,}",
        "PANEL_SELLERS": f"{audit['panel_sellers']:,}",
        "PANEL_MONTHS": str(audit["panel_months"]),
        "PANEL_START": audit["panel_start"],
        "PANEL_END": audit["panel_end"],
        "GMV_MILLION": f"{audit['gmv_total_brl'] / 1e6:.3f}",
        "ZERO_RATE": f"{audit['next_month_zero_rate']:.2%}",
        "DELIVERED_ORDERS": f"{audit['delivered_orders']:,}",
        "DELIVERED_ITEMS": f"{audit['delivered_items']:,}",
        "MISSING_MAX": f"{audit['missing_rate_max']:.2%}",
        "CANDIDATE_COUNT": str(len(CANDIDATE_FEATURES)),
        "ANALYSIS_COUNT": str(len(FEATURES)),
        "EXCLUDED_FEATURES": zh_list(list(EXCLUDED_FEATURES)),
        "STRICT_COUNT": str(len(strict)),
        "STRICT_LIST": zh_list(strict),
        "STABLE_COUNT": str(len(stable)),
        "STABLE_LIST": zh_list(stable),
        "Q10_COUNT": str(len(q10)),
        "Q10_LIST": zh_list(q10),
        "Q30_COUNT": str(len(q30)),
        "Q30_LIST": zh_list(q30),
        "PRETEST_COUNT": str(len(pretest)),
        "PRETEST_LIST": zh_list(pretest),
        "TRIMMED_LIST": zh_list(select_scenario("copula_trimmed_period")),
        "AOV_LIST": zh_list(
            scenarios.loc[
                scenarios["target"].eq("log_aov_next_month")
                & scenarios["selected_ebh"].astype(bool),
                "feature",
            ].tolist()
        ),
        "AI_LIST": zh_list(select_scenario("copula_ai_xgboost")),
        "CORE_LIST": zh_list(core),
        "SECONDARY_LIST": zh_list(secondary),
        "TOP_SHAP_5": zh_list(top_shap[:5]),
        "SIGNIFICANT_FE": zh_list(significant["feature"].tolist()),
        "BEST_MODEL": str(best["model"]),
        "BEST_RMSE": f"{best['rmse_log']:.3f}",
        "BEST_R2": f"{best['r2_log']:.3f}",
        "NAIVE_RMSE": f"{naive['rmse_log']:.3f}",
        "RMSE_IMPROVEMENT": f"{improvement:.1%}",
        "COPULA_MEAN_KS": f"{diag_copula['mean_marginal_ks']:.3f}",
        "COPULA_MAX_KS": f"{diag_copula['max_marginal_ks']:.3f}",
        "COPULA_COV_ERROR": f"{diag_copula['covariance_relative_error']:.3f}",
        "COPULA_CROSS_ASYM": f"{diag_copula['cross_covariance_asymmetry']:.3f}",
        "GAUSSIAN_MEAN_KS": f"{diag_gaussian['mean_marginal_ks']:.3f}",
        "SIM_FDP": f"{simulation['mean_fdp']:.3f}",
        "SIM_POWER": f"{simulation['mean_power']:.3f}",
        "SIM_DISCOVERIES": f"{simulation['mean_discoveries']:.2f}",
        "SIM_ABOVE_Q": f"{simulation['probability_fdp_above_q']:.1%}",
        "SIM_REPETITIONS": str(simulation["repetitions"]),
        "SIM_KNOCKOFF_REPETITIONS": str(simulation["knockoff_repetitions_per_dataset"]),
        "BASE_SIM_FDP": f"{simulation['single_knockoff_mean_fdp']:.3f}",
        "BASE_SIM_POWER": f"{simulation['single_knockoff_mean_power']:.3f}",
        "TRAIN_N": f"{repro['model_metadata']['train_n']:,}",
        "VALID_N": f"{repro['model_metadata']['validation_n']:,}",
        "TEST_N": f"{repro['model_metadata']['test_n']:,}",
        "PEAK_MONTH": str(peak["year_month"]),
        "PEAK_GMV_THOUSAND": f"{peak['gmv'] / 1e3:.1f}",
        "PANEL_SHA256": audit["panel_sha256"],
        "SEED": str(repro["seed"]),
        "PYTHON_VERSION": repro["python"].split()[0],
        "NUMPY_VERSION": repro["numpy"],
        "PANDAS_VERSION": repro["pandas"],
        "SCIPY_VERSION": repro["scipy"],
        "SKLEARN_VERSION": repro["scikit_learn"],
        "STATSMODELS_VERSION": repro["statsmodels"],
        "XGBOOST_VERSION": repro["xgboost"],
        "SOURCE_TABLE": markdown_table(
            source_table,
            ["数据域", "文件", "行数", "用途"],
            ["数据域", "文件", "实际行数", "用途"],
            {"行数": lambda value: f"{int(value):,}"},
        ),
        "FEATURE_TABLE": markdown_table(
            dictionary,
            ["编号", "feature", "label_zh", "状态"],
            ["编号", "变量名", "中文定义", "处理状态"],
        ),
        "KEY_DESC_TABLE": markdown_table(
            key_desc,
            ["label_zh", "n", "missing_pct", "mean", "std", "median", "spearman_next_gmv"],
            ["变量", "N", "缺失率", "均值", "标准差", "中位数", "Spearman相关"],
        ),
        "TOP_CORR_TABLE": markdown_table(
            top_corr,
            ["label_zh", "spearman_next_gmv"],
            ["变量", "与次月对数GMV的Spearman相关"],
        ),
        "PRIMARY_TABLE": markdown_table(
            primary_table,
            ["label_zh", "mean_w", "positive_w_rate", "frequency_fmt", "mean_evalue", "strict_fmt"],
            ["变量", "平均W", "W为正比例", "单轮入选频率", "平均e-value", "e-BH"],
        ),
        "SENSITIVITY_TABLE": markdown_table(
            sensitivity_table,
            ["alpha_eBH", "alpha_kn", "严格入选数", "严格入选变量"],
            ["最终FDR水平", "单轮水平", "严格入选数", "严格入选变量"],
        ),
        "METRIC_TABLE": markdown_table(
            metric_table,
            ["model", "n_test", "rmse_log", "mae_log", "r2_log", "wape_pct"],
            ["模型", "测试N", "RMSE(log)", "MAE(log)", "R2(log)", "WAPE"],
        ),
        "SHAP_TABLE": markdown_table(
            shap_table,
            ["label_zh", "mean_abs_shap", "mean_shap"],
            ["变量", "平均绝对SHAP", "平均SHAP"],
        ),
        "FIXED_TABLE": markdown_table(
            fixed_view,
            ["label_zh", "coefficient", "cluster_se", "p_value", "ci95_low", "ci95_high"],
            ["变量", "系数", "聚类SE", "p值", "95%CI下限", "95%CI上限"],
        ),
        "SCENARIO_TABLE": markdown_table(
            scenario_table,
            ["设定", "alpha_eBH", "alpha_kn", "严格入选数", "稳定候选数", "严格入选变量"],
            ["模型设定", "alpha_eBH", "alpha_kn", "严格入选", "稳定候选", "严格入选变量"],
        ),
        "APPENDIX_DESC_TABLE": markdown_table(
            appendix_desc,
            [
                "feature",
                "label_zh",
                "n",
                "missing_pct",
                "mean",
                "std",
                "median",
                "spearman_next_gmv",
            ],
            ["变量名", "中文含义", "N", "缺失率", "均值", "标准差", "中位数", "Spearman相关"],
        ),
    }
    return context


def build_markdown() -> str:
    text = TEMPLATE.read_text(encoding="utf-8")
    for key, value in build_context().items():
        text = text.replace("{{" + key + "}}", value)
    unresolved = sorted(set(re.findall(r"\{\{[A-Z0-9_]+\}\}", text)))
    if unresolved:
        raise ValueError(f"Unresolved template placeholders: {unresolved}")
    controls = [char for char in text if ord(char) < 32 and char not in "\n\t\r"]
    if controls:
        raise ValueError(f"Template contains control characters: {sorted(set(map(ord, controls)))}")
    return text


def make_reference_doc(path: Path) -> None:
    doc = Document()
    styles = doc.styles
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_after = Pt(0)

    title = doc.styles["Title"]
    title.font.name = "Times New Roman"
    title.font.size = Pt(22)
    title.font.bold = True
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.styles["Subtitle"]
    subtitle.font.size = Pt(16)
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if "Author" not in [style.name for style in styles]:
        author = styles.add_style("Author", WD_STYLE_TYPE.PARAGRAPH)
    else:
        author = styles["Author"]
    author.font.name = "Times New Roman"
    author.font.size = Pt(12)
    author._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    author.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author.paragraph_format.first_line_indent = Pt(0)

    for name, size, align in [
        ("Heading 1", 16, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 14, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.alignment = align
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.page_break_before = True

    caption = doc.styles["Caption"]
    caption.font.size = Pt(10.5)
    caption._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Pt(0)
    doc.save(path)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def cover_metadata_paragraph() -> OxmlElement:
    paragraph = OxmlElement("w:p")
    paragraph_properties = OxmlElement("w:pPr")
    justification = OxmlElement("w:jc")
    justification.set(qn("w:val"), "center")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:line"), "360")
    spacing.set(qn("w:lineRule"), "auto")
    paragraph_properties.extend([justification, spacing])
    paragraph.append(paragraph_properties)

    lines = [
        "学校：【待填写】",
        "学院：【待填写】",
        "专业：【待填写】",
        "研究生：【待填写】",
        "学号：【待填写】",
        "指导教师：【待填写】",
    ]
    for index, line in enumerate(lines):
        run = OxmlElement("w:r")
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), "宋体")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "24")
        run_properties.extend([fonts, size])
        run.append(run_properties)
        text = OxmlElement("w:t")
        text.text = line
        run.append(text)
        if index < len(lines) - 1:
            run.append(OxmlElement("w:br"))
        paragraph.append(run)
    return paragraph


def postprocess_docx(path: Path) -> None:
    doc = Document(path)
    doc.core_properties.title = "面向电商GMV监控的可控错误发现维度选择研究（v1）"
    doc.core_properties.subject = "Olist；Model-X Knockoff；e-value；可解释机器学习"
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(2.5)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.clear()
        add_page_field(footer)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings._element.append(update_fields)

    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "二〇二六年八月":
            paragraph._p.addprevious(cover_metadata_paragraph())
            break

    body = doc._element.body
    for index, child in enumerate(list(body)):
        has_toc_field = any(
            "TOC " in (node.text or "")
            for node in child.iter(qn("w:instrText"))
        )
        if child.tag == qn("w:sdt") and has_toc_field:
            for text_node in child.iter(qn("w:t")):
                if text_node.text == "Table of Contents":
                    text_node.text = "目录"
            page_break_paragraph = OxmlElement("w:p")
            run = OxmlElement("w:r")
            page_break = OxmlElement("w:br")
            page_break.set(qn("w:type"), "page")
            run.append(page_break)
            page_break_paragraph.append(run)
            body.insert(index, page_break_paragraph)
            break

    for table in doc.tables:
        table.style = "Table Grid"
        table.autofit = True
        for row_index, row in enumerate(table.rows):
            if row_index == 0:
                table_header = OxmlElement("w:tblHeader")
                table_header.set(qn("w:val"), "true")
                row._tr.get_or_add_trPr().append(table_header)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
                        run.font.name = "Times New Roman"
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for paragraph in doc.paragraphs:
        if re.match(r"^表\d", paragraph.text.strip()):
            paragraph.paragraph_format.keep_with_next = True
    doc.save(path)


def write_css(path: Path) -> None:
    path.write_text(
        """
@page {
  size: A4;
  margin: 25mm 25mm 25mm 30mm;
  @bottom-center { content: counter(page); font-size: 9pt; }
}
@page:first { @bottom-center { content: ""; } }
body { font-family: "Songti SC", "STSong", serif; font-size: 12pt; line-height: 1.75; color: #111; max-width: 160mm; margin: 0 auto; }
h1 { font-family: "Heiti SC", sans-serif; text-align: center; font-size: 18pt; page-break-before: always; margin-top: 0; }
h2 { font-family: "Heiti SC", sans-serif; font-size: 15pt; margin-top: 1.2em; break-after: avoid-page; }
h3 { font-family: "Heiti SC", sans-serif; font-size: 13pt; break-after: avoid-page; }
p { text-align: justify; text-indent: 2em; margin: 0.35em 0; }
p:has(+ table) { break-after: avoid-page; }
li p, td p, th p { text-indent: 0; }
a { color: #111; text-decoration: none; }
table { border-collapse: collapse; width: 100%; font-size: 8.5pt; margin: 0.8em 0; page-break-inside: auto; }
tr { page-break-inside: avoid; }
th, td { border: 1px solid #555; padding: 4px 6px; vertical-align: top; }
th { background: #f0f0f0; }
img { display: block; max-width: 95%; max-height: 210mm; margin: 1em auto; page-break-inside: avoid; }
.title { font-family: "Heiti SC", sans-serif; font-size: 24pt; text-align: center; margin-top: 52mm; }
.subtitle { text-align: center; font-size: 16pt; }
.author, .date { text-align: center; text-indent: 0; }
#TOC { page-break-before: always; page-break-after: always; }
#TOC::before { content: "目录"; display: block; font-family: "Heiti SC", sans-serif; font-size: 18pt; font-weight: bold; text-align: center; margin: 0 0 1.2em 0; }
""".strip(),
        encoding="utf-8",
    )


def run(command: list[str]) -> None:
    subprocess.run(command, check=True, cwd=V1_DIR)


def main() -> None:
    markdown_path = V1_DIR / "论文_v1.md"
    root_markdown_path = V1_DIR.parent / "论文_v1_最终版.md"
    docx_path = V1_DIR / "论文_v1.docx"
    html_path = V1_DIR / "论文_v1.html"
    pdf_path = V1_DIR / "论文_v1.pdf"
    reference_path = V1_DIR / "reference.docx"
    css_path = V1_DIR / "thesis.css"

    text = build_markdown()
    markdown_path.write_text(text, encoding="utf-8")
    root_text = text.replace("](figures/", "](v1/figures/")
    root_markdown_path.write_text(root_text, encoding="utf-8")
    make_reference_doc(reference_path)
    write_css(css_path)

    run(
        [
            "pandoc",
            str(markdown_path),
            "--from=markdown+tex_math_dollars",
            "--toc",
            "--toc-depth=3",
            f"--reference-doc={reference_path}",
            f"--resource-path={V1_DIR}",
            "-o",
            str(docx_path),
        ]
    )
    postprocess_docx(docx_path)
    run(
        [
            "pandoc",
            str(markdown_path),
            "--from=markdown+tex_math_dollars",
            "--standalone",
            "--toc",
            "--toc-depth=3",
            "--mathml",
            "--embed-resources",
            f"--css={css_path}",
            f"--resource-path={V1_DIR}",
            "-o",
            str(html_path),
        ]
    )

    chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
    if not chrome.exists():
        raise FileNotFoundError("Google Chrome is required for PDF generation")
    run(
        [
            str(chrome),
            "--headless",
            "--disable-gpu",
            "--allow-file-access-from-files",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path.resolve().as_uri(),
        ]
    )

    manifest = {
        "version": "v1",
        "markdown_characters": len(text),
        "chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", text)),
        "unresolved_placeholders": re.findall(r"\{\{[A-Z0-9_]+\}\}", text),
        "outputs": {
            path.name: path.stat().st_size
            for path in [markdown_path, root_markdown_path, docx_path, html_path, pdf_path]
        },
    }
    (RESULTS / "thesis_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
