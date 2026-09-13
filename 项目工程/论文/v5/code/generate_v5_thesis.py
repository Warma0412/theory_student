#!/usr/bin/env python3
"""Generate the complete V5 thesis with weekly-granularity robustness."""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd
from docx import Document


HERE = Path(__file__).resolve()
V5_DIR = HERE.parents[1]
PAPER_DIR = V5_DIR.parent
PROJECT_DIR = HERE.parents[3]
V1_DIR = PAPER_DIR / "v1"
V3_DIR = PAPER_DIR / "v3"
V4_DIR = PAPER_DIR / "v4"
RESULTS = V5_DIR / "results"
FIGURES = V5_DIR / "figures"
TEMPLATE = HERE.with_name("thesis_template_independent.md")

sys.path.insert(0, str(V3_DIR / "code"))
sys.path.insert(0, str(V4_DIR / "code"))
from generate_v3_thesis import (  # noqa: E402
    make_reference_doc,
    markdown_table,
    postprocess_docx,
    write_css,
)
from generate_v4_thesis import build_context as build_base_context  # noqa: E402


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def selected(frame: pd.DataFrame, q: float, method: str | None = None) -> pd.DataFrame:
    view = frame.loc[frame["q"].eq(q)].copy()
    if method is not None:
        view = view.loc[view["method"].eq(method)]
    return view.loc[view["selected_ebh"].astype(bool)].copy()


def list_zh(frame: pd.DataFrame) -> str:
    values = frame["label_zh"].tolist()
    return "、".join(values) if values else "无"


def fmt_percent(value: float, digits: int = 1) -> str:
    return f"{value:.{digits}%}"


def build_context() -> dict[str, str]:
    context = build_base_context()
    monthly_audit = load_json(V1_DIR / "results" / "data_audit.json")
    weekly_audit = load_json(RESULTS / "data_audit.json")
    metadata = load_json(RESULTS / "v5_weekly_reproducibility.json")
    diagnostics = pd.DataFrame(
        load_json(RESULTS / "weekly_knockoff_diagnostics.json")
    )
    primary = pd.read_csv(RESULTS / "weekly_knockoff_primary.csv")
    monthly_primary = pd.read_csv(V1_DIR / "results" / "knockoff_primary.csv")
    comparison = pd.read_csv(RESULTS / "monthly_weekly_comparison.csv")
    metrics = pd.read_csv(RESULTS / "weekly_predictive_metrics.csv")
    clustered = pd.read_csv(RESULTS / "clustered_knockoff_selection.csv")
    grouped = pd.read_csv(RESULTS / "group_knockoff_selection.csv")
    grip = pd.read_csv(RESULTS / "grip2_style_selection.csv")
    fixed_effects = pd.read_csv(RESULTS / "two_way_fixed_effects.csv")
    panel = pd.read_csv(RESULTS.parent / "data_processed" / "seller_week_panel.csv")

    weekly_q20 = selected(primary, 0.20)
    monthly_q20 = selected(monthly_primary, 0.20)
    comparison_q20 = comparison.loc[comparison["q"].eq(0.20)].copy()
    overlap = comparison_q20.loc[comparison_q20["selected_both"].astype(bool)].copy()
    monthly_only = comparison_q20.loc[
        comparison_q20["monthly_selected"].astype(bool)
        & ~comparison_q20["weekly_selected"].astype(bool)
    ].copy()
    weekly_only = comparison_q20.loc[
        comparison_q20["weekly_selected"].astype(bool)
        & ~comparison_q20["monthly_selected"].astype(bool)
    ].copy()
    weekly_stable = primary.loc[
        primary["q"].eq(0.20) & primary["selection_frequency"].ge(0.90)
    ].copy()
    stable_cross = overlap.loc[
        overlap["feature"].isin(weekly_stable["feature"])
    ].copy()

    grain_summary = pd.DataFrame(
        [
            {
                "口径": "月度主分析",
                "面板行": monthly_audit["panel_rows"],
                "卖家": monthly_audit["panel_sellers"],
                "期间数": monthly_audit["panel_months"],
                "目标零值率": fmt_percent(monthly_audit["next_month_zero_rate"]),
                "q=.20入选": len(monthly_q20),
                "研究角色": "确认性主结论",
            },
            {
                "口径": "周度稳健性",
                "面板行": weekly_audit["panel_rows"],
                "卖家": weekly_audit["panel_sellers"],
                "期间数": weekly_audit["panel_weeks"],
                "目标零值率": fmt_percent(weekly_audit["next_week_zero_rate"]),
                "q=.20入选": len(weekly_q20),
                "研究角色": "粒度稳健性补充",
            },
        ]
    )

    weekly_table = weekly_q20.sort_values(
        ["selection_frequency", "mean_w"], ascending=False
    ).copy()
    weekly_table["频率"] = weekly_table["selection_frequency"].map(fmt_percent)
    monthly_set = set(monthly_q20["feature"])
    weekly_table["月度同时入选"] = weekly_table["feature"].map(
        lambda feature: "是" if feature in monthly_set else "否"
    )
    weekly_table["平均W"] = weekly_table["mean_w"].map(lambda value: f"{value:.4f}")
    weekly_table["平均e值"] = weekly_table["mean_evalue"].map(
        lambda value: f"{value:.3f}"
    )

    overlap_table = comparison_q20.loc[
        comparison_q20["monthly_selected"].astype(bool)
        | comparison_q20["weekly_selected"].astype(bool)
    ].copy()
    overlap_table["证据"] = overlap_table.apply(
        lambda row: (
            "月周共同"
            if row["selected_both"]
            else ("仅月度" if row["monthly_selected"] else "仅周度")
        ),
        axis=1,
    )
    overlap_table["月频率"] = overlap_table["monthly_frequency"].map(fmt_percent)
    overlap_table["周频率"] = overlap_table["weekly_frequency"].map(fmt_percent)
    evidence_order = {"月周共同": 0, "仅月度": 1, "仅周度": 2}
    overlap_table["order"] = overlap_table["证据"].map(evidence_order)
    overlap_table = overlap_table.sort_values(
        ["order", "weekly_frequency", "monthly_frequency"],
        ascending=[True, False, False],
    )

    sensitivity_specs = [
        ("周度主分析", "weekly_knockoff_primary.csv", None),
        ("原始高斯生成", "weekly_knockoff_gaussian.csv", None),
        ("同周GMV对照", "weekly_same_week_sensitivity.csv", None),
        ("截尾时间窗", "weekly_trimmed_sensitivity.csv", None),
        ("下一周客单价", "weekly_aov_sensitivity.csv", None),
        ("测试期前样本", "weekly_pretest_knockoff.csv", None),
        ("XGBoost统计量", "weekly_xgboost_knockoff.csv", None),
    ]
    sensitivity_rows = []
    for name, filename, method in sensitivity_specs:
        frame = pd.read_csv(RESULTS / filename)
        view = selected(frame, 0.20, method)
        sensitivity_rows.append(
            {
                "情景": name,
                "重复数": metadata[
                    {
                        "周度主分析": "weekly_primary",
                        "原始高斯生成": "weekly_gaussian",
                        "同周GMV对照": "weekly_same_week",
                        "截尾时间窗": "weekly_trimmed",
                        "下一周客单价": "weekly_aov",
                        "测试期前样本": "weekly_pretest",
                        "XGBoost统计量": "weekly_xgboost_knockoff",
                    }[name]
                ]["repetitions"],
                "严格数": len(view),
                "与周度主集重合": len(
                    set(view["feature"]) & set(weekly_q20["feature"])
                ),
                "严格集合": list_zh(view),
            }
        )
    sensitivity_table = pd.DataFrame(sensitivity_rows)

    metric_view = metrics.copy()
    metric_view["RMSE"] = metric_view["rmse_log"].map(lambda value: f"{value:.3f}")
    metric_view["MAE"] = metric_view["mae_log"].map(lambda value: f"{value:.3f}")
    metric_view["R2"] = metric_view["r2_log"].map(lambda value: f"{value:.3f}")
    metric_view["WAPE"] = metric_view["wape_raw"].map(fmt_percent)
    metric_names = {
        "ExtraTrees": "Extra Trees",
        "XGBoost": "XGBoost",
        "XGBoost-weekly-FDR-set": "XGBoost（周度18项）",
        "Lasso": "Lasso",
        "Ridge": "Ridge",
        "MLP": "MLP",
        "Naive-current-GMV": "本周GMV朴素基线",
    }
    metric_view["模型"] = metric_view["model"].map(metric_names)

    cluster_rows = []
    for method, name in [
        ("clustered_between_seller", "卖家间均值层"),
        ("clustered_within_seller", "卖家内离差层"),
    ]:
        view = selected(clustered, 0.20, method)
        cluster_rows.append(
            {
                "层级": name,
                "样本量": metadata["clustered"][
                    method.replace("clustered_", "")
                ]["n"],
                "建模维度": metadata["clustered"][
                    method.replace("clustered_", "")
                ]["p"],
                "严格数": len(view),
                "严格集合": list_zh(view),
            }
        )
    cluster_table = pd.DataFrame(cluster_rows)

    group_q20 = selected(grouped, 0.20)
    grip_q20 = selected(grip, 0.20)
    grip_q30 = selected(grip, 0.30)
    fe_significant = fixed_effects.loc[fixed_effects["p_value"].lt(0.05)].copy()
    fe_significant["系数"] = fe_significant["coefficient"].map(
        lambda value: f"{value:.4f}"
    )
    fe_significant["p值"] = fe_significant["p_value"].map(
        lambda value: f"{value:.3g}"
    )

    copula = diagnostics.set_index("generator").loc["copula"]
    gaussian = diagnostics.set_index("generator").loc["gaussian"]
    best = metrics.sort_values("rmse_log").iloc[0]
    baseline = metrics.loc[metrics["model"].eq("Naive-current-GMV")].iloc[0]
    selected_model = metrics.loc[
        metrics["model"].eq("XGBoost-weekly-FDR-set")
    ].iloc[0]
    full_xgb = metrics.loc[metrics["model"].eq("XGBoost")].iloc[0]
    q20_comparison = metadata["monthly_weekly_comparison"]["0.2"]
    joint_diagnostics = pd.read_csv(
        V4_DIR / "results" / "v4_generator_diagnostic_summary.csv"
    )
    diagnostic_names = {
        "copula_mvr": "Copula-MVR",
        "v2_deep_train_calibrated": "基础深度生成器（训练集校准）",
        "v4_worst_swap_train_calibrated": "最坏swap深度生成器（训练集校准）",
    }
    joint_diagnostics["方法"] = joint_diagnostics["method"].map(
        diagnostic_names
    )
    joint_diagnostics["平均KS"] = joint_diagnostics["mean_marginal_ks"].map(
        lambda value: f"{value:.3f}"
    )
    joint_diagnostics["协方差误差"] = joint_diagnostics[
        "covariance_relative_error"
    ].map(lambda value: f"{value:.3f}")
    joint_diagnostics["深度核拒绝率"] = joint_diagnostics[
        "deep_kernel_rejection_rate_5pct"
    ].map(lambda value: f"{value:.0%}")
    joint_diagnostics["分类AUC"] = joint_diagnostics[
        "classifier_auc_mean"
    ].map(lambda value: f"{value:.3f}")

    references_text = (
        (V4_DIR / "论文_v4.md")
        .read_text(encoding="utf-8")
        .split("# 参考文献\n", maxsplit=1)[1]
        .split("\n# 附录A", maxsplit=1)[0]
        .strip()
    )
    reproducibility = pd.DataFrame(
        [
            ["随机种子", metadata["seed"]],
            ["月度面板SHA-256", monthly_audit["panel_sha256"]],
            ["周度面板SHA-256", weekly_audit["panel_sha256"]],
            ["Python", metadata["python"].split()[0]],
            ["NumPy", metadata["numpy"]],
            ["pandas", metadata["pandas"]],
            ["scikit-learn", metadata["sklearn"]],
            ["XGBoost", metadata["xgboost"]],
            ["PyTorch", metadata["torch"]],
            ["计算设备", metadata["device"]],
        ],
        columns=["项目", "记录值"],
    )
    shortlist = pd.read_csv(V3_DIR / "results" / "v3_shortlist.csv")
    shortlist["单轮频率"] = shortlist["selection_frequency"].map(fmt_percent)
    shortlist["贡献占比"] = shortlist[
        "contribution_share_within_fdr_set"
    ].map(fmt_percent)
    shortlist["累计贡献"] = shortlist["cumulative_contribution"].map(fmt_percent)
    k_path = pd.read_csv(V3_DIR / "results" / "v3_selection_path.csv")
    k_path = k_path.loc[
        k_path["k"].isin([4, 8, int(context["V3_DEFAULT_K"]), 14])
    ].copy()
    k_path["档位"] = k_path["k"].map(
        {
            4: "极简",
            8: "均衡",
            int(context["V3_DEFAULT_K"]): "默认推荐",
            14: "完整确认集",
        }
    )
    k_path["累计贡献"] = k_path["cumulative_contribution"].map(fmt_percent)
    k_path["验证RMSE"] = k_path["validation_rmse_log"].map(
        lambda value: f"{value:.3f}"
    )
    k_path["测试RMSE"] = k_path["rmse_log"].map(lambda value: f"{value:.3f}")

    context.update({
        "WEEK_ROWS": f"{weekly_audit['panel_rows']:,}",
        "WEEK_SELLERS": f"{weekly_audit['panel_sellers']:,}",
        "WEEK_COUNT": str(weekly_audit["panel_weeks"]),
        "WEEK_START": weekly_audit["panel_start"],
        "WEEK_END": weekly_audit["panel_end"],
        "WEEK_GMV": f"{weekly_audit['gmv_total_brl'] / 1e6:.3f}",
        "WEEK_ZERO_RATE": fmt_percent(weekly_audit["next_week_zero_rate"]),
        "WEEK_MEDIAN_ACTIVE": f"{panel.groupby('seller_id').size().median():.0f}",
        "WEEK_STRICT_COUNT": str(len(weekly_q20)),
        "WEEK_STRICT_LIST": list_zh(weekly_q20),
        "WEEK_STABLE_COUNT": str(len(weekly_stable)),
        "WEEK_STABLE_LIST": list_zh(weekly_stable),
        "OVERLAP_COUNT": str(len(overlap)),
        "OVERLAP_LIST": list_zh(overlap),
        "STABLE_CROSS_COUNT": str(len(stable_cross)),
        "STABLE_CROSS_LIST": list_zh(stable_cross),
        "MONTH_ONLY_COUNT": str(len(monthly_only)),
        "MONTH_ONLY_LIST": list_zh(monthly_only),
        "WEEK_ONLY_COUNT": str(len(weekly_only)),
        "WEEK_ONLY_LIST": list_zh(weekly_only),
        "JACCARD": f"{q20_comparison['jaccard']:.3f}",
        "WEEK_COPULA_KS": f"{copula['mean_marginal_ks']:.3f}",
        "WEEK_COPULA_MAX_KS": f"{copula['max_marginal_ks']:.3f}",
        "WEEK_COPULA_COV": f"{copula['covariance_relative_error']:.3f}",
        "WEEK_GAUSSIAN_KS": f"{gaussian['mean_marginal_ks']:.3f}",
        "WEEK_BEST_MODEL": metric_names[best["model"]],
        "WEEK_BEST_RMSE": f"{best['rmse_log']:.3f}",
        "WEEK_BEST_R2": f"{best['r2_log']:.3f}",
        "WEEK_RMSE_IMPROVEMENT": fmt_percent(
            1 - best["rmse_log"] / baseline["rmse_log"]
        ),
        "WEEK_SELECTED_RMSE": f"{selected_model['rmse_log']:.3f}",
        "WEEK_SELECTED_DELTA": fmt_percent(
            selected_model["rmse_log"] / full_xgb["rmse_log"] - 1,
            digits=2,
        ),
        "WEEK_BETWEEN_COUNT": str(len(selected(clustered, 0.20, "clustered_between_seller"))),
        "WEEK_WITHIN_COUNT": str(len(selected(clustered, 0.20, "clustered_within_seller"))),
        "WEEK_GROUP_COUNT": str(len(group_q20)),
        "WEEK_GROUP_LIST": list_zh(group_q20),
        "WEEK_GRIP_Q20_COUNT": str(len(grip_q20)),
        "WEEK_GRIP_Q30_COUNT": str(len(grip_q30)),
        "WEEK_GRIP_Q30_LIST": list_zh(grip_q30),
        "WEEK_FE_LIST": list_zh(fe_significant),
        "WEEK_PANEL_SHA": weekly_audit["panel_sha256"],
        "WEEK_GRAIN_TABLE": markdown_table(
            grain_summary,
            ["口径", "面板行", "卖家", "期间数", "目标零值率", "q=.20入选", "研究角色"],
            ["口径", "面板行", "卖家", "期间数", "目标零值率", "q=.20入选", "研究角色"],
        ),
        "WEEK_SELECTION_TABLE": markdown_table(
            weekly_table,
            ["label_zh", "平均W", "频率", "平均e值", "月度同时入选"],
            ["维度", "平均W", "单轮频率", "平均e-value", "月度同时入选"],
        ),
        "WEEK_OVERLAP_TABLE": markdown_table(
            overlap_table,
            ["label_zh", "证据", "月频率", "周频率"],
            ["维度", "粒度证据", "月度频率", "周度频率"],
        ),
        "WEEK_SENSITIVITY_TABLE": markdown_table(
            sensitivity_table,
            ["情景", "重复数", "严格数", "与周度主集重合"],
            ["情景", "重复数", "严格数", "与主集重合"],
        ),
        "WEEK_METRIC_TABLE": markdown_table(
            metric_view,
            ["模型", "n_test", "RMSE", "MAE", "R2", "WAPE"],
            ["模型", "测试N", "RMSE(log)", "MAE(log)", "R2", "WAPE"],
        ),
        "WEEK_CLUSTER_TABLE": markdown_table(
            cluster_table,
            ["层级", "样本量", "建模维度", "严格数", "严格集合"],
            ["层级", "N", "p", "严格数", "严格集合"],
        ),
        "WEEK_FE_TABLE": markdown_table(
            fe_significant,
            ["label_zh", "系数", "p值"],
            ["维度", "系数", "p值"],
        ),
        "JOINT_DIAGNOSTIC_TABLE": markdown_table(
            joint_diagnostics,
            ["方法", "平均KS", "协方差误差", "深度核拒绝率", "分类AUC"],
            ["方法", "平均KS", "协方差误差", "深度核拒绝率", "分类AUC"],
        ),
        "REPRODUCIBILITY_TABLE": markdown_table(
            reproducibility,
            ["项目", "记录值"],
            ["项目", "记录值"],
        ),
        "REFERENCES_TEXT": references_text,
        "K_SHORTLIST_COMPACT_TABLE": markdown_table(
            shortlist,
            [
                "shortlist_rank",
                "label_zh",
                "单轮频率",
                "贡献占比",
                "累计贡献",
            ],
            ["排序", "维度", "单轮频率", "贡献占比", "累计贡献"],
        ),
        "K_PATH_COMPACT_TABLE": markdown_table(
            k_path,
            ["档位", "k", "累计贡献", "验证RMSE", "测试RMSE"],
            ["档位", "K", "累计贡献", "验证RMSE", "测试RMSE"],
        ),
    })
    return context


def render(text: str, context: dict[str, str]) -> str:
    for key, value in context.items():
        text = text.replace("{{" + key + "}}", str(value))
    unresolved = sorted(set(re.findall(r"\{\{[A-Z0-9_]+\}\}", text)))
    if unresolved:
        raise ValueError(f"Unresolved V5 placeholders: {unresolved}")
    return text


def _build_legacy_markdown() -> str:
    context = build_context()
    text = (V4_DIR / "论文_v4.md").read_text(encoding="utf-8")
    text = text.replace(
        "基于Olist卖家月度面板的聚类与组级Knockoff、深度核诊断及参数化选择（v4）",
        "基于Olist卖家月度主分析与周度粒度稳健性的可控维度选择（v5）",
        1,
    )
    text = text.replace("date: 二〇二六年八月", "date: 二〇二六年九月", 1)
    text = text.replace(
        "本文档为研究生学位论文 v4 完整稿。",
        "本文档为研究生学位论文 v5 完整稿。",
        1,
    )
    text = text.replace(
        "V4新增结果、模型和日志保存在 `论文/v4`。",
        (
            "V4新增结果、模型和日志保存在 `论文/v4`。V5在不改写月度"
            "确认性主线的前提下，新增完整卖家-周面板计算；周度结果、每轮W、"
            "模型指标与日志保存在 `论文/v5`。"
        ),
        1,
    )

    abstract_addition = """
V5进一步检验结论对时间聚合粒度是否敏感。采用与月度主分析相同的已送达订单回溯口径，构建{{WEEK_ROWS}}个卖家-周观测、{{WEEK_SELLERS}}个卖家和{{WEEK_COUNT}}个周，使用本周31项维度预测下一周`log(1+GMV)`。周度目标零成交率为{{WEEK_ZERO_RATE}}，高于月度的25.86%。在相同60轮Copula-MVR、Lasso、`α_kn=0.10`和`α_eBH=0.20`下，周度严格入选{{WEEK_STRICT_COUNT}}项；与月度14项共同入选{{OVERLAP_COUNT}}项，Jaccard系数为{{JACCARD}}。跨粒度共同集合为{{OVERLAP_LIST}}。周度卖家间层入选{{WEEK_BETWEEN_COUNT}}项、卖家内层仍为{{WEEK_WITHIN_COUNT}}项。时间外测试最佳模型{{WEEK_BEST_MODEL}}的RMSE为{{WEEK_BEST_RMSE}}，周度18项XGBoost的RMSE为{{WEEK_SELECTED_RMSE}}。结果支持月度作为低频卖家主分析、周度作为更稀疏但信息量更大的粒度稳健性；周度18项不替换月度14项。
""".strip()
    keyword_marker = (
        "**关键词：** 电商 GMV；聚类数据 Knockoff；Group Knockoff；"
        "深度核 MMD；GRIP2；错误发现率"
    )
    text = text.replace(
        keyword_marker,
        abstract_addition + "\n\n" + keyword_marker + "；粒度稳健性",
        1,
    )

    english_addition = """
Version 5 adds a full seller-week robustness analysis without replacing the monthly confirmatory design. The weekly panel contains {{WEEK_ROWS}} observations for {{WEEK_SELLERS}} sellers over {{WEEK_COUNT}} weeks, with a {{WEEK_ZERO_RATE}} zero rate for next-week GMV. At the same final e-BH level of 0.20, the weekly analysis selects {{WEEK_STRICT_COUNT}} variables; {{OVERLAP_COUNT}} overlap with the 14 monthly discoveries, yielding a Jaccard index of {{JACCARD}}. The between-seller weekly layer selects {{WEEK_BETWEEN_COUNT}} variables, whereas the within-seller layer remains empty. The weekly results are therefore interpreted as granularity robustness rather than a replacement confirmatory set.
""".strip()
    key_marker = (
        "**Key words:** E-commerce GMV; clustered knockoffs; group knockoffs; "
        "deep-kernel MMD; trajectory importance; false discovery rate"
    )
    text = text.replace(
        key_marker,
        english_addition + "\n\n" + key_marker + "; granularity robustness",
        1,
    )

    method_addition = """
## 4.19 V5周度粒度稳健性设计

V5保留卖家-月为确认性主分析，并使用卖家-周作为粒度稳健性。周定义为周一至周日，特征周覆盖{{WEEK_START}}至{{WEEK_END}}，共{{WEEK_COUNT}}周。分析单位仍要求卖家在当前周至少有一笔已送达订单；若下一周无已送达订单，则下一周GMV记为零。周度目标定义为

$$
Y_{s,w+1}=\\log(1+\\mathrm{GMV}_{s,w+1}).
$$

周度分析完全复用31项候选维度、随机化秩高斯Copula、Ledoit-Wolf收缩协方差、MVR Knockoff、Lasso竞争统计量、60轮e-value聚合以及`α_kn=0.10`、`α_eBH=0.20`。同时重跑原始高斯、同周GMV、截尾时间窗、下一周客单价、测试期前样本和XGBoost统计量，并重跑卖家均值/离差分层、十二业务组、GRIP2式轨迹统计、双向固定效应和自然时间外预测。

预测样本按自然周切分：2017-01-02至2017-12-25为训练期，2018-01-01至2018-04-30为验证期，2018-05-07至2018-07-30为测试期。周度结果用于回答“结论是否依赖月度聚合”，不重新定义V1的主假设，也不改变V3基于月度14项确认池建立的参数K路径。

需要特别说明，V5沿用“最终已送达订单、按购买周归属”的回溯口径，以确保月度与周度可比。物流和评价字段在对应周末未必已经完全可见，因此该实验是粒度稳健性检验，不是严格的在线部署回放。真实部署应按事件可见时间归属或至少滞后一周。

""".strip()
    text = text.replace(
        "# 第5章 实证结果",
        method_addition + "\n\n# 第5章 实证结果",
        1,
    )

    result_addition = """
## 5.17 V5周度粒度稳健性结果

### 5.17.1 面板规模、稀疏性与生成诊断

周度面板覆盖{{WEEK_START}}至{{WEEK_END}}，包含{{WEEK_ROWS}}个卖家-周、{{WEEK_SELLERS}}个卖家和{{WEEK_COUNT}}周，GMV合计{{WEEK_GMV}}百万雷亚尔。每个卖家活跃周数中位数仅为{{WEEK_MEDIAN_ACTIVE}}，下一周GMV为零的比例为{{WEEK_ZERO_RATE}}，明显高于月度25.86%。周度增加了横截面行数和时间点，却同时提高零膨胀、单订单周和共线性，不能简单解释为“样本更多所以更优”。

**表5-16 月度主分析与周度稳健性口径**

{{WEEK_GRAIN_TABLE}}

![图5-13 Olist周度GMV](figures/fig_v5_weekly_gmv.png)

Copula生成器平均边际KS为{{WEEK_COPULA_KS}}、最大KS为{{WEEK_COPULA_MAX_KS}}、协方差相对误差为{{WEEK_COPULA_COV}}；原始高斯平均边际KS为{{WEEK_GAUSSIAN_KS}}。尽管周度高度共线设计在MVR矩阵计算时产生浮点警告，全部W矩阵和结果数值均为有限值，且Copula边际匹配明显优于高斯基线，因此正式周度分支继续采用Copula-MVR。

### 5.17.2 周度主选择与月周交集

周度主分析在0.10水平无严格发现；在0.20水平严格入选{{WEEK_STRICT_COUNT}}项：{{WEEK_STRICT_LIST}}。其中{{WEEK_STABLE_COUNT}}项的单轮频率至少为90%：{{WEEK_STABLE_LIST}}。

**表5-17 周度q=0.20严格入选变量**

{{WEEK_SELECTION_TABLE}}

![图5-14 周度Knockoff单轮入选频率](figures/fig_v5_weekly_selection.png)

月度14项与周度18项共同入选{{OVERLAP_COUNT}}项，Jaccard系数为{{JACCARD}}。共同集合为{{OVERLAP_LIST}}；其中又满足周度单轮频率至少90%的跨粒度高稳定核心有{{STABLE_CROSS_COUNT}}项：{{STABLE_CROSS_LIST}}。月度特有{{MONTH_ONLY_COUNT}}项为{{MONTH_ONLY_LIST}}；周度特有{{WEEK_ONLY_COUNT}}项为{{WEEK_ONLY_LIST}}。

**表5-18 月度与周度全部入选变量对照**

{{WEEK_OVERLAP_TABLE}}

![图5-15 月度与周度入选频率](figures/fig_v5_monthly_weekly_frequency.png)

交集主要覆盖交易规模、价格、承运准备、商品标题和市场覆盖，说明这些信息不依赖单一聚合粒度。月度特有的送达时长、预计提前量和描述长度更像慢变化信号；周度特有的商品物理属性、支付结构与营销标记则更容易受单订单周构成波动影响。后者只能作为周度补充，不应因出现于18项集合就自动升级为主看板指标。

### 5.17.3 周度敏感性、聚类与组级结果

**表5-19 周度Knockoff敏感性**

{{WEEK_SENSITIVITY_TABLE}}

同周GMV对照选出29项，远多于下一周主分析，直接展示了价格、件数和订单数参与GMV定义造成的机械关系。测试期前样本选出14项，与完整周度主集重合13项；XGBoost统计量选出16项，与主集重合15项；下一周客单价情景选出15项且全部属于周度主集。这些结果支持主要信号具有一定模型和时间稳健性，但并非所有周度特有项均稳定。

**表5-20 周度聚类分层结果**

{{WEEK_CLUSTER_TABLE}}

周度卖家间层在0.20水平入选{{WEEK_BETWEEN_COUNT}}项，卖家内层为{{WEEK_WITHIN_COUNT}}项。十二业务组在0.20水平入选{{WEEK_GROUP_COUNT}}组：{{WEEK_GROUP_LIST}}。GRIP2式统计量在0.20水平为{{WEEK_GRIP_Q20_COUNT}}项，在0.30水平为{{WEEK_GRIP_Q30_COUNT}}项：{{WEEK_GRIP_Q30_LIST}}。与月度V4相同，卖家内层仍无严格发现，说明提高到周频并未解决同一卖家短期动态信号不足的问题。

### 5.17.4 周度时间外预测与固定效应

**表5-21 周度时间外预测表现**

{{WEEK_METRIC_TABLE}}

周度最佳模型为{{WEEK_BEST_MODEL}}，测试RMSE为{{WEEK_BEST_RMSE}}、$R^2$为{{WEEK_BEST_R2}}，相对本周GMV朴素基线改善{{WEEK_RMSE_IMPROVEMENT}}。只使用周度18项的XGBoost测试RMSE为{{WEEK_SELECTED_RMSE}}，相对全31维XGBoost变化{{WEEK_SELECTED_DELTA}}，说明压缩后几乎不损失预测性能。

![图5-16 周度时间外预测比较](figures/fig_v5_weekly_prediction.png)

周度双向固定效应在5%水平显著的变量为{{WEEK_FE_LIST}}：

{{WEEK_FE_TABLE}}

该结果仍是控制卖家和周效应后的条件关联，不是因果效应。尤其履约字段按最终已送达订单回溯计算，不能直接视为周末实时可行动信息。

### 5.17.5 粒度选择结论

V5支持“月度主分析、周度稳健性、日度不作为主线”的设计。月度零值更少、每个单元聚合订单更多，更适合低频卖家的稳定确认；周度提供更多时间点并能发现短期构成变化，但零膨胀和共线性更强。正式论文继续以月度14项和默认K=10为主结论，周度18项作为敏感性集合，月周共同11项作为跨粒度核心证据。若未来取得曝光、库存、广告等可实时观测变量及更长序列，可再建立真正的周度在线预测版本。

""".strip()
    text = text.replace(
        "# 第6章 业务应用与讨论",
        result_addition + "\n\n# 第6章 业务应用与讨论",
        1,
    )

    discussion_addition = """
## 6.8 月度主看板与周度预警的协同

月度与周度不是二选一。月度层负责统计确认、季度复核和参数K展示；周度层负责观察方向变化和短期异常。建议将月周共同11项标记为跨粒度核心，其中周度高稳定的10项优先用于预警。月度特有项保留在履约与内容诊断页，周度特有项先进入观察池，只有在连续滚动窗口、测试期前筛选和实时可见口径下重复出现，才考虑进入常驻看板。

周度系统上线时应重构时间语义：支付信息按支付发生时间，发货信息按承运时间，评价信息按评价创建时间进入特征；当前V5的最终送达回溯口径只回答统计粒度问题。该区分避免把论文的回溯性稳健结论误写成可直接部署的在线预测能力。

""".strip()
    text = text.replace(
        "# 第7章 结论与展望",
        discussion_addition + "\n\n# 第7章 结论与展望",
        1,
    )

    conclusion_marker = (
        "V4进一步发现：整表信号主要由卖家间差异驱动，卖家内层在0.20水平无严格发现；"
        "组级0.20水平受低组数离散门槛限制；最坏swap生成器仍未通过深度核与分类器准入；"
        "GRIP2式重要性在0.30水平恢复5项但主水平仍为0。V4没有改变14项确认集和默认K=10，"
        "而是显著收紧其适用边界。"
    )
    conclusion_addition = """

V5在完整周度面板上复现同一统计流程。33,190个卖家-周的下一周零成交率为40.46%，在0.20水平选出18项，与月度14项交集11项、Jaccard为0.524；卖家内层仍为0项。周度18项模型的测试RMSE为2.421，与全31维XGBoost的2.420几乎一致。综合偏差、方差、业务节奏和跨粒度结果，月度继续承担确认性主结论，周度作为稳健性和短期预警证据。
""".rstrip()
    text = text.replace(
        conclusion_marker,
        conclusion_marker + conclusion_addition,
        1,
    )
    text = text.replace(
        "第七，固定效应、SHAP和参数K均不能替代因果实验。",
        (
            "第七，固定效应、SHAP和参数K均不能替代因果实验。第八，V5周度"
            "分析沿用最终送达回溯口径，不能视为严格的实时部署回放。"
        ),
        1,
    )

    appendix = """
# 附录G V5周度稳健性复现摘要

| 项目 | 记录值 |
|---|---|
| 周度面板 | {{WEEK_ROWS}}个卖家-周，{{WEEK_SELLERS}}个卖家，{{WEEK_COUNT}}周 |
| 特征期 | {{WEEK_START}}至{{WEEK_END}} |
| 下一周零成交率 | {{WEEK_ZERO_RATE}} |
| 建模维度 | 31 |
| 主Knockoff重复 | 60 |
| 最终/单轮水平 | 0.20 / 0.10 |
| 周度严格集合 | {{WEEK_STRICT_COUNT}}项 |
| 月周交集 | {{OVERLAP_COUNT}}项，Jaccard={{JACCARD}} |
| 聚类分层 | 卖家间{{WEEK_BETWEEN_COUNT}}项；卖家内{{WEEK_WITHIN_COUNT}}项 |
| Group Knockoff | q=.20为{{WEEK_GROUP_COUNT}}组 |
| GRIP2式 | q=.20为{{WEEK_GRIP_Q20_COUNT}}项；q=.30为{{WEEK_GRIP_Q30_COUNT}}项 |
| 周度最佳模型 | {{WEEK_BEST_MODEL}}，RMSE={{WEEK_BEST_RMSE}} |
| 面板SHA-256 | {{WEEK_PANEL_SHA}} |

V5运行命令：

```bash
论文/v1/.venv/bin/python 论文/v5/code/build_weekly_panel.py
论文/v1/.venv/bin/python 论文/v5/code/run_v5_weekly_analysis.py
论文/v1/.venv/bin/python 论文/v5/code/generate_v5_thesis.py
```

全部周度W矩阵、选择表、预测结果、诊断和日志保存在`论文/v5/results`与`论文/v5/logs`。

""".strip()
    text = text.replace(
        "# 附录G v4提交前检查清单",
        appendix + "\n\n# 附录H v5提交前检查清单",
        1,
    )
    text = text.replace("V4运行命令：", "V4运行命令：", 1)
    text = text.replace(
        "V4没有推翻14项确认集",
        "V4没有推翻14项确认集",
        1,
    )
    return render(text, context)


def build_markdown() -> str:
    """Render the standalone thesis without development-version narrative."""
    text = TEMPLATE.read_text(encoding="utf-8")
    return render(text, build_context())


def copy_base_figures() -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    for source in (V4_DIR / "figures").glob("*.png"):
        destination = FIGURES / source.name
        if not destination.exists() or destination.stat().st_mtime < source.stat().st_mtime:
            shutil.copy2(source, destination)
    aliases = {
        "fig_v2_generator_diagnostics.png": "fig08_deep_generator_diagnostics.png",
        "fig_v2_swap_mmd.png": "fig09_swap_mmd.png",
        "fig_v2_predictive_models.png": "fig10_predictive_models.png",
        "fig_v2_deep_selection_comparison.png": "fig11_deep_selection.png",
        "fig_v4_cluster_group.png": "fig12_cluster_group.png",
        "fig_v4_exchangeability.png": "fig13_exchangeability.png",
        "fig_v4_grip_importance.png": "fig14_grip_importance.png",
        "fig_v5_weekly_gmv.png": "fig15_weekly_gmv.png",
        "fig_v5_weekly_selection.png": "fig16_weekly_selection.png",
        "fig_v5_monthly_weekly_frequency.png": "fig17_monthly_weekly_frequency.png",
        "fig_v5_weekly_prediction.png": "fig18_weekly_prediction.png",
    }
    for source_name, destination_name in aliases.items():
        source = FIGURES / source_name
        destination = FIGURES / destination_name
        if source.exists() and not destination.exists():
            shutil.copy2(source, destination)


def run(command: list[str]) -> None:
    subprocess.run(command, check=True, cwd=V5_DIR)


def main() -> None:
    copy_base_figures()
    text = build_markdown()
    markdown_path = V5_DIR / "论文_v5.md"
    root_markdown_path = PAPER_DIR / "论文_v5_最终版.md"
    docx_path = V5_DIR / "论文_v5.docx"
    html_path = V5_DIR / "论文_v5.html"
    pdf_path = V5_DIR / "论文_v5.pdf"
    reference_path = V5_DIR / "reference.docx"
    css_path = V5_DIR / "thesis.css"

    markdown_path.write_text(text, encoding="utf-8")
    root_markdown_path.write_text(
        text.replace("](figures/", "](v5/figures/"),
        encoding="utf-8",
    )
    make_reference_doc(reference_path)
    write_css(css_path)
    run(
        [
            "pandoc",
            str(markdown_path),
            "--from=markdown+tex_math_dollars",
            "--toc",
            "--toc-depth=3",
            f"--reference-doc={reference_path}",
            f"--resource-path={V5_DIR}",
            "-o",
            str(docx_path),
        ]
    )
    postprocess_docx(docx_path)
    document = Document(docx_path)
    document.core_properties.title = (
        "基于去随机化Model-X Knockoff的电商GMV多粒度受控维度选择研究"
    )
    document.core_properties.subject = (
        "月度主分析；周度粒度稳健性；Model-X Knockoff；e-BH"
    )
    document.save(docx_path)
    run(
        [
            "pandoc",
            str(markdown_path),
            "--from=markdown+tex_math_dollars",
            "--standalone",
            "--toc",
            "--toc-depth=3",
            "--mathml",
            "--embed-resources",
            f"--css={css_path}",
            f"--resource-path={V5_DIR}",
            "-o",
            str(html_path),
        ]
    )
    chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
    run(
        [
            str(chrome),
            "--headless",
            "--disable-gpu",
            "--allow-file-access-from-files",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path.resolve().as_uri(),
        ]
    )
    manifest = {
        "document_role": "standalone_thesis",
        "markdown_characters": len(text),
        "chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", text)),
        "unresolved_placeholders": [],
        "outputs": {
            path.name: path.stat().st_size
            for path in [
                markdown_path,
                root_markdown_path,
                docx_path,
                html_path,
                pdf_path,
            ]
        },
    }
    (RESULTS / "thesis_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
