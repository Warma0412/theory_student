#!/usr/bin/env python3
"""Validate V5 weekly results and all document deliverables."""

from __future__ import annotations

import json
import re
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document


HERE = Path(__file__).resolve()
V5_DIR = HERE.parents[1]
PAPER_DIR = V5_DIR.parent
RESULTS = V5_DIR / "results"


def pdf_info(path: Path) -> tuple[int, str]:
    output = subprocess.run(
        ["pdfinfo", str(path)],
        check=True,
        capture_output=True,
        text=True,
    ).stdout
    page_match = re.search(r"^Pages:\s+(\d+)$", output, re.MULTILINE)
    size_match = re.search(r"^Page size:\s+(.+)$", output, re.MULTILINE)
    if not page_match or not size_match:
        raise ValueError(f"Unable to read PDF metadata: {path}")
    return int(page_match.group(1)), size_match.group(1)


def pdf_text(path: Path) -> str:
    return subprocess.run(
        ["pdftotext", str(path), "-"],
        check=True,
        capture_output=True,
        text=True,
    ).stdout


def selection_count(
    frame: pd.DataFrame, q: float, method: str | None = None
) -> int:
    view = frame.loc[frame["q"].eq(q)]
    if method is not None:
        view = view.loc[view["method"].eq(method)]
    return int(view["selected_ebh"].astype(bool).sum())


def assert_finite_csv(path: Path) -> None:
    frame = pd.read_csv(path)
    numeric = frame.select_dtypes(include=[np.number])
    values = numeric.to_numpy()
    assert np.isfinite(values[~np.isnan(values)]).all(), path


def main() -> None:
    documents = [
        V5_DIR / "论文_v5.md",
        V5_DIR / "论文_v5.docx",
        V5_DIR / "论文_v5.html",
        V5_DIR / "论文_v5.pdf",
        PAPER_DIR / "论文_v5_最终版.md",
        PAPER_DIR / "论文_v5_精简版.pdf",
        PAPER_DIR / "论文_v5_极限精简版.pdf",
    ]
    assert all(path.exists() and path.stat().st_size > 0 for path in documents)

    for path in RESULTS.glob("*.csv"):
        assert_finite_csv(path)

    primary = pd.read_csv(RESULTS / "weekly_knockoff_primary.csv")
    comparison = pd.read_csv(RESULTS / "monthly_weekly_comparison.csv")
    clustered = pd.read_csv(RESULTS / "clustered_knockoff_selection.csv")
    grouped = pd.read_csv(RESULTS / "group_knockoff_selection.csv")
    grip = pd.read_csv(RESULTS / "grip2_style_selection.csv")
    metrics = pd.read_csv(RESULTS / "weekly_predictive_metrics.csv")
    diagnostics = pd.DataFrame(
        json.loads(
            (RESULTS / "weekly_knockoff_diagnostics.json").read_text(
                encoding="utf-8"
            )
        )
    )
    metadata = json.loads(
        (RESULTS / "v5_weekly_reproducibility.json").read_text(
            encoding="utf-8"
        )
    )

    assert primary.shape == (93, 13)
    assert selection_count(primary, 0.10) == 0
    assert selection_count(primary, 0.20) == 18
    assert selection_count(primary, 0.30) == 18
    assert int(
        primary.loc[
            primary["q"].eq(0.20), "selection_frequency"
        ].ge(0.90).sum()
    ) == 13

    comparison_q20 = comparison.loc[comparison["q"].eq(0.20)]
    monthly_count = int(comparison_q20["monthly_selected"].astype(bool).sum())
    weekly_count = int(comparison_q20["weekly_selected"].astype(bool).sum())
    intersection_count = int(comparison_q20["selected_both"].astype(bool).sum())
    union_count = int(
        (
            comparison_q20["monthly_selected"].astype(bool)
            | comparison_q20["weekly_selected"].astype(bool)
        ).sum()
    )
    jaccard = intersection_count / union_count
    assert (monthly_count, weekly_count, intersection_count) == (14, 18, 11)
    assert abs(jaccard - 11 / 21) < 1e-12

    assert selection_count(
        clustered, 0.20, "clustered_between_seller"
    ) == 12
    assert selection_count(
        clustered, 0.20, "clustered_within_seller"
    ) == 0
    assert selection_count(grouped, 0.20) == 10
    assert selection_count(grip, 0.20) == 0
    assert selection_count(grip, 0.30) == 12

    same_week = pd.read_csv(RESULTS / "weekly_same_week_sensitivity.csv")
    pretest = pd.read_csv(RESULTS / "weekly_pretest_knockoff.csv")
    xgboost_w = pd.read_csv(RESULTS / "weekly_xgboost_knockoff.csv")
    assert selection_count(same_week, 0.20) == 29
    assert selection_count(pretest, 0.20) == 14
    assert selection_count(xgboost_w, 0.20) == 16

    best = metrics.sort_values("rmse_log").iloc[0]
    selected_model = metrics.loc[
        metrics["model"].eq("XGBoost-weekly-FDR-set")
    ].iloc[0]
    assert best["model"] == "ExtraTrees"
    assert abs(best["rmse_log"] - 2.416023606812565) < 1e-9
    assert abs(selected_model["rmse_log"] - 2.420783431130472) < 1e-9

    diagnostics = diagnostics.set_index("generator")
    assert diagnostics.loc["copula", "mean_marginal_ks"] < 0.02
    assert (
        diagnostics.loc["copula", "mean_marginal_ks"]
        < diagnostics.loc["gaussian", "mean_marginal_ks"]
    )
    assert metadata["role"] == (
        "weekly_granularity_robustness_not_replacement_of_monthly_primary"
    )
    assert metadata["grip2_style_antisymmetry"]["passed_tolerance_1e_5"]
    assert metadata["panel_rows"] == 33190
    assert metadata["panel_weeks"] == 83

    markdown = (V5_DIR / "论文_v5.md").read_text(encoding="utf-8")
    assert not re.findall(r"\{\{[A-Z0-9_]+\}\}", markdown)
    for phrase in [
        "面向电商GMV监控的多粒度可控维度选择研究",
        "多粒度可控维度选择方法",
        "周度粒度稳健性",
        "讨论与经营应用",
        "不是严格的在线部署回放",
        "Jaccard系数为0.524",
        "本文识别的是条件预测信息",
    ]:
        assert phrase in markdown
    forbidden_patterns = [
        r"\bV[1-5]\b",
        r"Version\s+[1-5]",
        r"上一版",
        r"前一版",
        r"版本沿革",
        r"保留V[1-5]",
        r"新增V[1-5]",
    ]
    for pattern in forbidden_patterns:
        assert not re.search(pattern, markdown, flags=re.IGNORECASE), pattern
    abstract_zh = markdown.split("# 摘要\n", maxsplit=1)[1].split(
        "**关键词：**", maxsplit=1
    )[0]
    abstract_zh_chars = len(re.findall(r"[\u4e00-\u9fff]", abstract_zh))
    assert 500 <= abstract_zh_chars <= 1000
    references = markdown.split("# 参考文献\n", maxsplit=1)[1].split(
        "\n# 附录A", maxsplit=1
    )[0]
    reference_numbers = [
        int(value) for value in re.findall(r"^\[(\d+)\]", references, re.MULTILINE)
    ]
    assert reference_numbers == list(range(1, 46))
    chapter_section_counts = {}
    for number, content in re.findall(
        r"^# 第(\d+)章[^\n]*\n(.*?)(?=^# 第\d+章|^# 参考文献)",
        markdown,
        flags=re.MULTILINE | re.DOTALL,
    ):
        chapter_section_counts[number] = len(
            re.findall(r"^## ", content, re.MULTILINE)
        )
    assert chapter_section_counts == {
        "1": 5,
        "2": 5,
        "3": 5,
        "4": 5,
        "5": 5,
        "6": 5,
        "7": 4,
    }

    docx = Document(V5_DIR / "论文_v5.docx")
    assert len(docx.paragraphs) > 250
    assert len(docx.tables) >= 20
    assert len(docx.inline_shapes) >= 10
    assert docx.core_properties.title == (
        "面向电商GMV监控的多粒度可控维度选择研究"
    )

    full_pages, full_size = pdf_info(V5_DIR / "论文_v5.pdf")
    condensed_pages, condensed_size = pdf_info(
        PAPER_DIR / "论文_v5_精简版.pdf"
    )
    ultra_pages, ultra_size = pdf_info(
        PAPER_DIR / "论文_v5_极限精简版.pdf"
    )
    assert full_pages >= 45
    assert 12 <= condensed_pages <= 25
    assert ultra_pages == 5
    assert all(
        "A4" in size for size in [full_size, condensed_size, ultra_size]
    )
    for path in [
        V5_DIR / "论文_v5.pdf",
        PAPER_DIR / "论文_v5_精简版.pdf",
        PAPER_DIR / "论文_v5_极限精简版.pdf",
    ]:
        text = pdf_text(path)
        for pattern in forbidden_patterns:
            assert not re.search(pattern, text, flags=re.IGNORECASE), (
                path,
                pattern,
            )

    report = {
        "status": "passed",
        "weekly_primary": {
            "q10_count": 0,
            "q20_count": 18,
            "q30_count": 18,
            "stable_90pct_count": 13,
        },
        "monthly_weekly": {
            "monthly_count": monthly_count,
            "weekly_count": weekly_count,
            "intersection_count": intersection_count,
            "jaccard": jaccard,
        },
        "clustered_q20": {"between": 12, "within": 0},
        "group_q20_count": 10,
        "grip2": {"q20_count": 0, "q30_count": 12},
        "predictive": {
            "best_model": best["model"],
            "best_rmse": float(best["rmse_log"]),
            "selected_set_rmse": float(selected_model["rmse_log"]),
        },
        "pdf_pages": {
            "complete": full_pages,
            "condensed": condensed_pages,
            "ultra": ultra_pages,
        },
        "thesis_structure": {
            "chinese_abstract_characters": abstract_zh_chars,
            "reference_count": len(reference_numbers),
            "chapter_section_counts": chapter_section_counts,
        },
        "standalone_thesis_version_markers": 0,
        "all_result_csv_numeric_values_finite": True,
    }
    (RESULTS / "validation_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
