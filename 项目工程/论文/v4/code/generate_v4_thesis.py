#!/usr/bin/env python3
"""Generate the complete V4 thesis from V1-V4 machine-readable results."""

from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import pandas as pd
from docx import Document


HERE = Path(__file__).resolve()
V4_DIR = HERE.parents[1]
PROJECT_DIR = HERE.parents[3]
V1_DIR = PROJECT_DIR / "论文" / "v1"
V2_DIR = PROJECT_DIR / "论文" / "v2"
V3_DIR = PROJECT_DIR / "论文" / "v3"
V3_CODE = V3_DIR / "code"
RESULTS = V4_DIR / "results"
TEMPLATE = HERE.with_name("thesis_template_v4.md")
sys.path.insert(0, str(V3_CODE))

from generate_v3_thesis import (  # noqa: E402
    build_context as build_v3_context,
    labels,
    make_reference_doc,
    markdown_table,
    postprocess_docx,
    write_css,
)


def load_json(name: str):
    return json.loads((RESULTS / name).read_text(encoding="utf-8"))


def selected(frame: pd.DataFrame, q: float, method: str | None = None) -> list[str]:
    view = frame.loc[frame["q"].eq(q)]
    if method is not None:
        view = view.loc[view["method"].eq(method)]
    return view.loc[view["selected_ebh"].astype(bool), "feature"].tolist()


def selected_labels(
    frame: pd.DataFrame, q: float, method: str | None = None
) -> list[str]:
    view = frame.loc[frame["q"].eq(q)]
    if method is not None:
        view = view.loc[view["method"].eq(method)]
    return view.loc[view["selected_ebh"].astype(bool), "label_zh"].tolist()


def build_context() -> dict[str, str]:
    context = build_v3_context()
    clustered = pd.read_csv(RESULTS / "clustered_knockoff_selection.csv")
    grouped = pd.read_csv(RESULTS / "group_knockoff_selection.csv")
    diagnostics = pd.read_csv(RESULTS / "v4_generator_diagnostic_summary.csv")
    detail = pd.read_csv(RESULTS / "v4_exchangeability_diagnostics.csv")
    deep_selection = pd.read_csv(
        RESULTS / "v4_adversarial_generator_selection.csv"
    )
    grip = pd.read_csv(RESULTS / "grip2_style_selection.csv")
    reproducibility = load_json("v4_reproducibility.json")

    between_method = "clustered_between_seller"
    within_method = "clustered_within_seller"
    between = selected(clustered, 0.20, between_method)
    within = selected(clustered, 0.20, within_method)
    group_q20 = selected_labels(grouped, 0.20)
    group_q30 = selected_labels(grouped, 0.30)
    deep_formal = selected(deep_selection, 0.20)
    grip_q20 = selected(grip, 0.20)
    grip_q30 = selected(grip, 0.30)

    cluster_rows = []
    for method, name in [
        (between_method, "卖家间均值层"),
        (within_method, "卖家内离差层"),
    ]:
        view = clustered.loc[
            clustered["method"].eq(method) & clustered["q"].eq(0.20)
        ]
        strict_labels = view.loc[
            view["selected_ebh"].astype(bool), "label_zh"
        ].tolist()
        stable_labels = view.loc[
            view["selection_frequency"].ge(0.90), "label_zh"
        ].tolist()
        cluster_rows.append(
            {
                "层级": name,
                "样本量": reproducibility["clustered"][
                    method.replace("clustered_", "")
                ]["n"],
                "建模维度": reproducibility["clustered"][
                    method.replace("clustered_", "")
                ]["p"],
                "严格数": len(strict_labels),
                "严格集合": "、".join(strict_labels) if strict_labels else "无",
                "稳定数": len(stable_labels),
            }
        )
    cluster_table = pd.DataFrame(cluster_rows)

    group_rows = []
    for q in (0.20, 0.30):
        view = grouped.loc[grouped["q"].eq(q)]
        strict_labels = view.loc[
            view["selected_ebh"].astype(bool), "label_zh"
        ].tolist()
        stable_labels = view.loc[
            view["selection_frequency"].ge(0.90), "label_zh"
        ].tolist()
        group_rows.append(
            {
                "最终水平": q,
                "单轮水平": q / 2,
                "严格组数": len(strict_labels),
                "严格业务组": "、".join(strict_labels) if strict_labels else "无",
                "高稳定组": "、".join(stable_labels) if stable_labels else "无",
            }
        )
    group_table = pd.DataFrame(group_rows)

    diagnostic_names = {
        "copula_mvr": "Copula-MVR",
        "v2_deep_train_calibrated": "V2深度生成（训练集校准）",
        "v4_worst_swap_train_calibrated": "V4最坏swap（训练集校准）",
    }
    diagnostic_view = diagnostics.copy()
    diagnostic_view["方法"] = diagnostic_view["method"].map(diagnostic_names)
    diagnostic_view["deep_reject_fmt"] = diagnostic_view[
        "deep_kernel_rejection_rate_5pct"
    ].map(lambda value: f"{value:.0%}")
    diagnostic_view["classifier_reject_fmt"] = diagnostic_view[
        "classifier_rejection_rate_5pct"
    ].map(lambda value: f"{value:.0%}")

    deep_index = diagnostics.set_index("method")
    v2_row = deep_index.loc["v2_deep_train_calibrated"]
    v4_row = deep_index.loc["v4_worst_swap_train_calibrated"]
    deep_valid = bool(
        reproducibility["v4_generator_passed_prespecified_diagnostics"]
    )
    deep_decision = "通过，可进入推断" if deep_valid else "未通过，不作FDR声明"

    grip_q20_view = (
        grip.loc[grip["q"].eq(0.20)]
        .sort_values(["mean_w", "selection_frequency"], ascending=False)
        .head(15)
        .copy()
    )
    grip_q20_view["frequency_fmt"] = grip_q20_view[
        "selection_frequency"
    ].map(lambda value: f"{value:.1%}")
    q30_selected_set = set(selected(grip, 0.30))
    grip_q20_view["selected_q30"] = grip_q20_view["feature"].map(
        lambda feature: "是" if feature in q30_selected_set else "否"
    )

    repro = pd.DataFrame(
        [
            ["聚类分层重复", reproducibility["config"]["clustered_repetitions"]],
            ["组级Knockoff重复", reproducibility["config"]["group_repetitions"]],
            ["最坏swap生成重复", reproducibility["config"]["generator_repetitions"]],
            ["GRIP2式重复", reproducibility["config"]["grip_repetitions"]],
            ["深度核置换次数", reproducibility["diagnostic_permutations"]],
            ["深度核最小p值", reproducibility["minimum_attainable_p"]],
            ["分类器置换次数", reproducibility["classifier_permutations"]],
            ["分类器最小p值", reproducibility["classifier_minimum_attainable_p"]],
            ["边际校准拟合范围", reproducibility["marginal_calibration_fit_scope"]],
            ["V4深度生成准入", deep_decision],
            [
                "GRIP2式反对称",
                "通过"
                if reproducibility["grip2_style_antisymmetry"][
                    "passed_tolerance_1e_5"
                ]
                else "未通过",
            ],
            ["面板SHA-256", reproducibility["panel_sha256"]],
        ],
        columns=["项目", "记录值"],
    )

    context.update(
        {
            "V4_BETWEEN_COUNT": str(len(between)),
            "V4_BETWEEN_LIST": labels(between),
            "V4_WITHIN_COUNT": str(len(within)),
            "V4_WITHIN_LIST": labels(within),
            "V4_GROUP_Q20_COUNT": str(len(group_q20)),
            "V4_GROUP_Q20_LIST": "、".join(group_q20) if group_q20 else "无",
            "V4_GROUP_Q30_COUNT": str(len(group_q30)),
            "V4_GROUP_Q30_LIST": "、".join(group_q30) if group_q30 else "无",
            "V4_V2_TRAIN_KS": f"{v2_row['mean_marginal_ks']:.3f}",
            "V4_V2_DEEP_REJECT": (
                f"{v2_row['deep_kernel_rejection_rate_5pct']:.0%}"
            ),
            "V4_V2_AUC": f"{v2_row['classifier_auc_mean']:.3f}",
            "V4_DEEP_KS": f"{v4_row['mean_marginal_ks']:.3f}",
            "V4_DEEP_COV": f"{v4_row['covariance_relative_error']:.3f}",
            "V4_DEEP_KERNEL_REJECT": (
                f"{v4_row['deep_kernel_rejection_rate_5pct']:.0%}"
            ),
            "V4_DEEP_AUC": f"{v4_row['classifier_auc_mean']:.3f}",
            "V4_DEEP_DECISION": deep_decision,
            "V4_DEEP_FORMAL_LIST": labels(deep_formal),
            "V4_GRIP_ERROR": (
                f"{reproducibility['grip2_style_antisymmetry']['max_absolute_error']:.3e}"
            ),
            "V4_GRIP_Q20_COUNT": str(len(grip_q20)),
            "V4_GRIP_Q20_LIST": labels(grip_q20),
            "V4_GRIP_Q30_COUNT": str(len(grip_q30)),
            "V4_GRIP_Q30_LIST": labels(grip_q30),
            "V4_CLUSTER_TABLE": markdown_table(
                cluster_table,
                ["层级", "样本量", "建模维度", "严格数", "稳定数", "严格集合"],
                ["层级", "N", "p", "严格数", "稳定数", "严格集合"],
            ),
            "V4_GROUP_TABLE": markdown_table(
                group_table,
                ["最终水平", "单轮水平", "严格组数", "严格业务组", "高稳定组"],
                ["最终水平", "单轮水平", "严格组数", "严格业务组", "高稳定组"],
            ),
            "V4_DIAGNOSTIC_TABLE": markdown_table(
                diagnostic_view,
                [
                    "方法",
                    "mean_marginal_ks",
                    "covariance_relative_error",
                    "cross_covariance_asymmetry",
                    "mean_real_knockoff_correlation",
                    "deep_kernel_rejection_rate_5pct",
                    "classifier_auc_mean",
                ],
                [
                    "方法",
                    "平均KS",
                    "协方差误差",
                    "交叉非对称",
                    "真伪相关",
                    "深度核拒绝率",
                    "分类AUC",
                ],
            ),
            "V4_GRIP_TABLE": markdown_table(
                grip_q20_view,
                [
                    "label_zh",
                    "mean_w",
                    "frequency_fmt",
                    "mean_evalue",
                    "selected_q30",
                ],
                ["维度", "平均W", "0.20单轮频率", "0.20平均e-value", "0.30入选"],
            ),
            "V4_REPRO_TABLE": markdown_table(
                repro, ["项目", "记录值"], ["项目", "记录值"]
            ),
        }
    )
    return context


def build_markdown() -> str:
    text = TEMPLATE.read_text(encoding="utf-8")
    for key, value in build_context().items():
        text = text.replace("{{" + key + "}}", str(value))
    unresolved = sorted(set(re.findall(r"\{\{[A-Z0-9_]+\}\}", text)))
    if unresolved:
        raise ValueError(f"Unresolved V4 placeholders: {unresolved}")
    controls = [
        character
        for character in text
        if ord(character) < 32 and character not in "\n\t\r"
    ]
    if controls:
        raise ValueError("V4 thesis contains control characters")
    return text


def run(command: list[str]) -> None:
    subprocess.run(command, check=True, cwd=V4_DIR)


def main() -> None:
    markdown_path = V4_DIR / "论文_v4.md"
    root_markdown_path = V4_DIR.parent / "论文_v4_最终版.md"
    docx_path = V4_DIR / "论文_v4.docx"
    html_path = V4_DIR / "论文_v4.html"
    pdf_path = V4_DIR / "论文_v4.pdf"
    reference_path = V4_DIR / "reference.docx"
    css_path = V4_DIR / "thesis.css"
    text = build_markdown()
    markdown_path.write_text(text, encoding="utf-8")
    root_markdown_path.write_text(
        text.replace("](figures/", "](v4/figures/"), encoding="utf-8"
    )
    make_reference_doc(reference_path)
    write_css(css_path)
    run(
        [
            "pandoc",
            str(markdown_path),
            "--from=markdown+tex_math_dollars",
            "--toc",
            "--toc-depth=3",
            f"--reference-doc={reference_path}",
            f"--resource-path={V4_DIR}",
            "-o",
            str(docx_path),
        ]
    )
    postprocess_docx(docx_path)
    doc = Document(docx_path)
    doc.core_properties.title = (
        "面向电商GMV监控的可控错误发现维度选择研究（v4）"
    )
    doc.core_properties.subject = (
        "聚类Knockoff；Group Knockoff；深度核MMD；轨迹聚合重要性"
    )
    doc.save(docx_path)
    run(
        [
            "pandoc",
            str(markdown_path),
            "--from=markdown+tex_math_dollars",
            "--standalone",
            "--toc",
            "--toc-depth=3",
            "--mathml",
            "--embed-resources",
            f"--css={css_path}",
            f"--resource-path={V4_DIR}",
            "-o",
            str(html_path),
        ]
    )
    chrome = Path("/Applications/Google Chrome.app/Contents/MacOS/Google Chrome")
    run(
        [
            str(chrome),
            "--headless",
            "--disable-gpu",
            "--allow-file-access-from-files",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_path}",
            html_path.resolve().as_uri(),
        ]
    )
    manifest = {
        "version": "v4",
        "markdown_characters": len(text),
        "chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", text)),
        "unresolved_placeholders": [],
        "outputs": {
            path.name: path.stat().st_size
            for path in [
                markdown_path,
                root_markdown_path,
                docx_path,
                html_path,
                pdf_path,
            ]
        },
    }
    (RESULTS / "thesis_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
